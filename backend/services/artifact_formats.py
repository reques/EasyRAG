"""Office generation and bounded, passive previews of generated artifacts."""
import csv
import io
import json
import math
import re
from html import escape
from itertools import islice
from pathlib import Path
from zipfile import ZipFile

MAX_CELLS = 100_000
PREVIEW_ROWS = 200
PREVIEW_COLUMNS = 50


def document_blocks(content):
    """Parse Markdown into document blocks shared by Word and PDF writers."""
    import markdown
    from lxml import html

    root = html.fragment_fromstring(markdown.markdown(content, extensions=['tables', 'fenced_code']), create_parent='div')
    for unwanted in root.xpath('.//script|.//style|.//iframe|.//object'):
        unwanted.drop_tree()
    return root


def write_docx(content):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    def add_inline(paragraph, node, bold=None, italic=None):
        bold = True if node.tag in {'strong', 'b'} else bold
        italic = True if node.tag in {'em', 'i'} else italic
        if node.text:
            run = paragraph.add_run(node.text)
            run.bold, run.italic = bold, italic
        for child in node:
            if child.tag == 'br':
                paragraph.add_run().add_break()
            else:
                add_inline(paragraph, child, bold, italic)
            if child.tail:
                run = paragraph.add_run(child.tail)
                run.bold, run.italic = bold, italic

    def add_block(node):
        if node.tag == 'table':
            rows = node.xpath('./thead/tr|./tbody/tr|./tr')
            width = max((len(row) for row in rows), default=0)
            if not width:
                return
            table = document.add_table(rows=0, cols=width)
            table.style = 'Table Grid'
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    add_inline(cell.paragraphs[0], value, bold=value.tag == 'th')
            return
        if node.tag in {'ul', 'ol'}:
            for item in node:
                paragraph = document.add_paragraph(style='List Bullet' if node.tag == 'ul' else 'List Number')
                add_inline(paragraph, item)
            return
        if re.fullmatch(r'h[1-6]', str(node.tag)):
            paragraph = document.add_heading(level=int(node.tag[1]))
        else:
            paragraph = document.add_paragraph()
        add_inline(paragraph, node)

    for node in document_blocks(content):
        add_block(node)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def workbook_spec(content):
    """Accept JSON sheets; use CSV as a convenient single-sheet input."""
    if content.lstrip().startswith(('{', '[')):
        spec = json.loads(content)
        sheets = spec.get('sheets') if isinstance(spec, dict) else [{'name': 'Sheet1', 'rows': spec}]
    else:
        sheets = [{'name': 'Sheet1', 'rows': list(csv.reader(io.StringIO(content)))}]
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 20:
        raise ValueError('Excel requires 1–20 sheets: {"sheets":[{"name":"Sheet1","rows":[["Name","Value"],["A",1]]}]}')
    names, total = set(), 0
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValueError('Each sheet must contain name and rows')
        name, rows = sheet.get('name', 'Sheet1'), sheet.get('rows')
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or re.search(r'[\\/*?:\[\]\x00-\x1f]', name) or name.casefold() in names:
            raise ValueError('Sheet names must be unique, 1–31 characters, without \\ / * ? : [ ]')
        names.add(name.casefold())
        if not isinstance(rows, list) or not rows or len(rows) > 10000:
            raise ValueError('Each sheet requires 1–10000 rows')
        for row in rows:
            if not isinstance(row, list) or len(row) > 100:
                raise ValueError('Each row must be an array with at most 100 cells')
            total += len(row)
            if total > MAX_CELLS:
                raise ValueError('Workbook exceeds 100000 cells')
            for value in row:
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ValueError('Cells must be strings, numbers, booleans or null')
                if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)):
                    raise ValueError('Invalid Excel cell text or text longer than 32767 characters')
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError('Excel numbers must be finite')
    return sheets


def write_xlsx(content):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in workbook_spec(content):
        sheet = workbook.create_sheet(spec.get('name', 'Sheet1'))
        widths = {}
        for row_index, row in enumerate(spec['rows'], 1):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_index, column, value)
                # Text is literal. Never turn data beginning with '=' into executable formulas.
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                widths[column] = min(45, max(widths.get(column, 12), len(str(value or '')) + 2))
                if row_index == 1:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='297F91')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def write_pdf(content):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

    font_path = next((path for path in (
        '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc',
        'C:/Windows/Fonts/msyh.ttc', 'C:/Windows/Fonts/simsun.ttc',
    ) if Path(path).is_file()), None)
    font = 'EasyRAG-CJK' if font_path else 'STSong-Light'
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, font_path, subfontIndex=0) if font_path else UnicodeCIDFont(font))
    normal = ParagraphStyle('body', fontName=font, fontSize=11, leading=17, wordWrap='CJK', spaceAfter=8)
    output, story = io.BytesIO(), []
    def paragraph(text, style=normal):
        return Paragraph(escape(text).replace('\n', '<br/>'), style)
    for node in document_blocks(content):
        if node.tag == 'table':
            rows = [[paragraph(''.join(cell.itertext())) for cell in row] for row in node.xpath('./thead/tr|./tbody/tr|./tr')]
            if rows:
                width = max(map(len, rows))
                rows = [row + [''] * (width - len(row)) for row in rows]
                table = LongTable(rows, colWidths=[475 / width] * width, repeatRows=1, splitInRow=1)
                table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.lightgrey), ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
                story.extend([table, Spacer(1, 10)])
        elif node.tag in {'ul', 'ol'}:
            story.extend(paragraph(f'{i}. ' + ''.join(item.itertext())) for i, item in enumerate(node, 1))
        else:
            style = normal
            if re.fullmatch(r'h[1-6]', str(node.tag)):
                size = 22 - int(node.tag[1]) * 2
                style = ParagraphStyle('heading', parent=normal, fontSize=size, leading=size + 6, spaceBefore=10)
            story.append(paragraph(''.join(node.itertext()), style))
    SimpleDocTemplate(output, leftMargin=60, rightMargin=60).build(story or [paragraph(' ')])
    return output.getvalue()


def validate_archive(data):
    # A compressed Office file must not expand without bound during preview.
    with ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
            raise ValueError('Office file is too large to preview')


def preview_file(filename, data):
    extension = filename.rsplit('.', 1)[-1].lower()
    if extension == 'docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        validate_archive(data)
        document = Document(io.BytesIO(data))
        blocks = []
        for element in document.element.body:
            if element.tag.endswith('}p'):
                p = Paragraph(element, document)
                style = p.style.name if p.style else ''
                tag = 'h' + style[-1] if re.fullmatch(r'Heading [1-6]', style) else 'p'
                text = ''.join(('<strong>' if r.bold else '') + ('<em>' if r.italic else '') + escape(r.text).replace('\n', '<br>') + ('</em>' if r.italic else '') + ('</strong>' if r.bold else '') for r in p.runs)
                blocks.append(f'<{tag}>{text}</{tag}>')
            elif element.tag.endswith('}tbl'):
                table = Table(element, document)
                blocks.append('<table>' + ''.join('<tr>' + ''.join('<td>' + escape(c.text) + '</td>' for c in row.cells) + '</tr>' for row in table.rows) + '</table>')
        return {'kind': 'html', 'content': ''.join(blocks), 'note': '内容预览，分页与 Word 中的排版可能不同'}
    if extension == 'xlsx':
        from openpyxl import load_workbook
        validate_archive(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = []
            for sheet in workbook.worksheets[:20]:
                rows = [[value if isinstance(value, (str, int, float, bool)) or value is None else str(value) for value in row]
                        for row in sheet.iter_rows(max_row=min(sheet.max_row or 0, PREVIEW_ROWS), max_col=min(sheet.max_column or 0, PREVIEW_COLUMNS), values_only=True)]
                sheets.append({'name': sheet.title, 'rows': rows, 'total_rows': sheet.max_row,
                               'truncated': (sheet.max_row or 0) > PREVIEW_ROWS or (sheet.max_column or 0) > PREVIEW_COLUMNS})
            return {'kind': 'sheets', 'sheets': sheets}
        finally:
            workbook.close()
    text = data.decode('utf-8-sig')
    if extension == 'csv':
        rows = list(islice(csv.reader(io.StringIO(text)), PREVIEW_ROWS + 1))
        truncated = len(rows) > PREVIEW_ROWS or any(len(row) > PREVIEW_COLUMNS for row in rows)
        return {'kind': 'sheets', 'sheets': [{'name': 'CSV', 'rows': [row[:PREVIEW_COLUMNS] for row in rows[:PREVIEW_ROWS]], 'truncated': truncated}]}
    if extension == 'md':
        import markdown
        return {'kind': 'html', 'content': markdown.markdown(text, extensions=['tables', 'fenced_code'])}
    if extension == 'html':
        return {'kind': 'html', 'content': text}
    if extension == 'json':
        text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
    return {'kind': 'text', 'content': text}
