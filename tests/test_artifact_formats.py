import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.services.artifact_service import prepare_file, normalize_filename
from backend.services.artifact_formats import preview_file, workbook_spec


def test_word_is_real_office_document_with_headings_tables_and_emphasis():
    data, mime = prepare_file('报告.docx', '# 分析报告\n\n这是 **重要内容**。\n\n- 第一点\n- 第二点\n\n| 名称 | 金额 |\n| --- | --- |\n| 项目甲 | 100 |')
    document = Document(io.BytesIO(data))
    assert data.startswith(b'PK')
    assert 'wordprocessingml' in mime
    assert document.paragraphs[0].style.name == 'Heading 1'
    assert document.paragraphs[0].text == '分析报告'
    assert any(r.bold and r.text == '重要内容' for p in document.paragraphs for r in p.runs)
    assert document.tables[0].cell(1, 0).text == '项目甲'
    preview = preview_file('报告.docx', data)
    assert preview['kind'] == 'html'
    assert '<h1>分析报告</h1>' in preview['content']
    assert '<strong>重要内容</strong>' in preview['content']
    assert '项目甲' in preview['content']


def test_excel_multiple_sheets_and_literal_formula_text():
    content = json.dumps({'sheets': [
        {'name': '销售', 'rows': [['产品', '金额', '备注'], ['甲', 100, '=HYPERLINK("https://example.com")'], ['乙', 0, False]]},
        {'name': '汇总', 'rows': [['总额'], [100]]},
    ]})
    data, mime = prepare_file('销售.xlsx', content)
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert workbook.sheetnames == ['销售', '汇总']
    assert workbook['销售']['B2'].value == 100
    assert workbook['销售']['C2'].data_type == 's'
    assert workbook['销售']['B3'].value == 0
    assert workbook['销售'].freeze_panes == 'A2'
    assert 'spreadsheetml' in mime
    preview = preview_file('销售.xlsx', data)
    assert [sheet['name'] for sheet in preview['sheets']] == workbook.sheetnames
    assert preview['sheets'][0]['rows'][2] == ['乙', 0, False]
    workbook.close()


def test_csv_input_to_excel_and_preview_limits():
    data, _ = prepare_file('table.xlsx', 'name,value\n' + '\n'.join(f'行{i},{i}' for i in range(220)))
    sheet = preview_file('table.xlsx', data)['sheets'][0]
    assert sheet['truncated']
    assert len(sheet['rows']) == 200
    assert sheet['total_rows'] == 221
    assert sheet['rows'][1] == ['行0', '0']


@pytest.mark.parametrize('spec', [
    {}, {'sheets': []}, {'sheets': [{'name': 'bad/name', 'rows': [[1]]}]},
    {'sheets': [{'name': 'X', 'rows': [[1]]}, {'name': 'x', 'rows': [[2]]}]},
    {'sheets': [{'name': 'X', 'rows': [[{'formula': '=1+1'}]]}]},
    {'sheets': [{'name': 'X', 'rows': [[float('inf')]]}]},
    {'sheets': [{'name': 'X', 'rows': [['a' * 32768]]}]},
])
def test_reject_invalid_workbooks(spec):
    with pytest.raises(ValueError):
        workbook_spec(json.dumps(spec))


def test_pdf_is_readable_and_contains_chinese():
    data, mime = prepare_file('报告.pdf', '# 分析报告\n\n测试中文内容和 English 123。\n\n| 名称 | 金额 |\n| --- | --- |\n| 甲 | 100 |')
    assert data.startswith(b'%PDF-') and mime == 'application/pdf'
    reader = PdfReader(io.BytesIO(data))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert '分析报告' in text
    assert '100' in text


def test_legacy_extensions_generate_modern_files():
    assert normalize_filename('报告.DOC') == '报告.docx'
    assert normalize_filename('数据.xls') == '数据.xlsx'
    assert prepare_file('报告.doc', '# 标题')[0].startswith(b'PK')
    assert prepare_file('数据.xls', 'a,b\n1,2')[0].startswith(b'PK')


@pytest.mark.parametrize('filename,content,kind', [
    ('a.md', '# 标题', 'html'), ('a.html', '<h1>标题</h1>', 'html'),
    ('a.csv', '名称,金额\n甲,100', 'sheets'), ('a.txt', '文本', 'text'),
    ('a.json', '{"a": 1}', 'text'),
])
def test_original_formats_preview(filename, content, kind):
    assert preview_file(filename, content.encode())['kind'] == kind
